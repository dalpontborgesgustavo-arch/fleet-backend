"""Read the October 2026 production workbook as values; never evaluates formulas or changes it."""
import datetime as dt
import json
import sys
from decimal import Decimal, ROUND_HALF_UP

from openpyxl import load_workbook


def number(value, places):
    if value is None or value == "":
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        raise ValueError(f"Unexpected numeric input: {value!r}")
    decimal = Decimal(str(value)).quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)
    if decimal < 0:
        raise ValueError(f"Negative production value: {value!r}")
    return format(decimal, "f")


def text(value):
    value = str(value).strip() if value is not None else ""
    return value if value and value != "0" else None


def main(path):
    book = load_workbook(path, data_only=True, read_only=True)
    output = []
    for sheet in book.worksheets:
        if sheet.title.strip().startswith("Resumo"):
            continue
        header_row = next((index for index in range(1, min(sheet.max_row, 20) + 1)
                           if sheet.cell(index, 5).value == "OBRA"
                           and isinstance(sheet.cell(index, 9).value, dt.datetime)
                           and sheet.cell(index, 9).value.strftime("%Y-%m") == "2026-10"), None)
        if header_row is None:
            raise ValueError(f"Expected October 2026 daily grid in {sheet.title}")
        for excel_row, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            work = text(row[4] if len(row) > 4 else None)
            activity = text(row[5] if len(row) > 5 else None)
            unit = text(row[6] if len(row) > 6 else None)
            if not (activity and unit):
                continue
            daily = []
            for offset in range(31):
                value = number(row[8 + offset], 3)
                if value is not None:
                    daily.append({"date": f"2026-10-{offset + 1:02d}", "quantity": value})
            weeks = []
            for index in range(5):
                # Excel SUM treats blank plan cells as zero in this historical workbook.
                weekly_plan = number(row[46 + index], 3) or "0"  # AU:AY
                month_plan = number(row[52 + index], 3) or "0"  # BA:BE
                weeks.append({"index": index, "weeklyPlanQuantity": weekly_plan, "monthPlanQuantity": month_plan})
            if not work and not (any(Decimal(day["quantity"]) for day in daily)
                                 or any(Decimal(week["weeklyPlanQuantity"]) or Decimal(week["monthPlanQuantity"]) for week in weeks)):
                continue
            output.append({
                "competence": "2026-10",
                "engineer": sheet.title.strip(),
                "municipality": text(row[2]),
                "workCode": text(row[3]),
                "workName": work or "Obra não informada na planilha",
                "sourceWorkNameMissing": work is None,
                "activity": activity,
                "unit": unit,
                "unitPrice": number(row[58], 15),  # BG
                "sourceSheet": sheet.title,
                "sourceExcelRow": excel_row,
                "daily": daily,
                "weeks": weeks,
            })
    if not output:
        raise ValueError("No production lines found")
    # Compare source values with the workbook's cached JR summary before allowing any write.
    summary_sheet = next((sheet for sheet in book.worksheets if sheet.title.strip().startswith("Resumo")), None)
    if summary_sheet is None:
        raise ValueError("JR summary sheet missing")
    ranges = [(1, 3), (4, 10), (11, 17), (18, 24), (25, 31)]
    mismatches = []
    for summary_row in range(5, 13):
        activity = text(summary_sheet.cell(summary_row, 1).value)
        unit = text(summary_sheet.cell(summary_row, 2).value)
        if not activity or not unit:
            continue
        matching = [line for line in output if line["activity"] == activity and line["unit"] == unit]
        for week_index, (start_day, end_day) in enumerate(ranges):
            column = 3 + 6 * week_index
            actual = sum((Decimal(day["quantity"]) for line in matching for day in line["daily"]
                          if start_day <= int(day["date"][-2:]) <= end_day), Decimal(0))
            monthly = sum((Decimal(line["weeks"][week_index]["monthPlanQuantity"]) for line in matching), Decimal(0))
            weekly = sum((Decimal(line["weeks"][week_index]["weeklyPlanQuantity"]) for line in matching), Decimal(0))
            value = sum((Decimal(day["quantity"]) * Decimal(line["unitPrice"] or "0")
                         for line in matching for day in line["daily"]
                         if start_day <= int(day["date"][-2:]) <= end_day), Decimal(0))
            for label, calculated, offset, tolerance in [
                ("monthly", monthly, 0, Decimal("0.2")),
                ("weekly", weekly, 1, Decimal("0.2")),
                ("executed", actual, 2, Decimal("0.2")),
                ("value", value, 3, Decimal("0.05")),
            ]:
                cached = summary_sheet.cell(summary_row, column + offset).value
                if cached is None or isinstance(cached, bool) or not isinstance(cached, (int, float)):
                    continue
                if abs(calculated - Decimal(str(cached))) > tolerance:
                    mismatches.append(f"{activity}/{unit} week {week_index + 1} {label}: {calculated} != {cached}")
    if mismatches:
        raise ValueError("Workbook summary mismatch; import halted: " + "; ".join(mismatches[:12]))
    json.dump(output, sys.stdout, ensure_ascii=False)


if __name__ == "__main__":
    main(sys.argv[1])
