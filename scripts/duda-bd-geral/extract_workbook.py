"""Read-only, reproducible BD Geral extract for private migration staging.

Never saves the workbook. The resulting JSONL contains company financial data:
keep it out of Git, verify its hashes, and transfer only through approved JR paths.
"""

import argparse
import datetime as dt
import hashlib
import json
import pathlib
import sys
import warnings

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
from openpyxl.utils import get_column_letter, range_boundaries


HEADERS_ROW = 6
FIRST_BODY_ROW = 7
LAST_COLUMN = 29  # AC: the Excel table's end, not worksheet max_column.
EXCLUDED_COLUMNS = {8, 9}  # H: contract value; I: contract balance (user decision).
# Read-only SERVER BI comparison on 21/10/2026: these rows have business-value
# differences from the conflicted copy even after removing H/I. Never use the
# main copy as a certified financial source until Duda reconciles them.
AUDITED_CONFLICT_ROWS = {1380, 1413, 1418, 1438, 1439, 1456, 1510}


def serializable(value):
    if isinstance(value, ArrayFormula):
        return {"kind": "arrayFormula", "text": value.text, "ref": value.ref}
    if isinstance(value, DataTableFormula):
        return {"kind": "dataTableFormula", "ref": value.ref}
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, bytes):
        return value.hex()
    return value


def month_value(value):
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime("%Y-%m")
    if isinstance(value, str):
        try:
            return dt.datetime.fromisoformat(value).strftime("%Y-%m")
        except ValueError:
            return None
    return None


def text_value(value):
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def canonical_json(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def extract(source: pathlib.Path, output_dir: pathlib.Path):
    source_sha = hashlib.sha256(source.read_bytes()).hexdigest()
    # A newer workbook is still provisional until the disputed values are
    # reconciled with Duda; a changed file hash is not proof of resolution.
    conflict_rows = AUDITED_CONFLICT_ROWS
    output_dir.mkdir(parents=True, exist_ok=True)
    warnings.filterwarnings("ignore", message="Unknown extension is not supported")
    formulas = openpyxl.load_workbook(source, read_only=False, data_only=False)
    cached = openpyxl.load_workbook(source, read_only=False, data_only=True)
    ws = formulas["BD Geral"]
    values_ws = cached["BD Geral"]
    table = ws.tables.get("Tabela1")
    if table is None:
        raise ValueError("BD Geral source table is missing")
    first_col, first_row, last_col, last_body_row = range_boundaries(table.ref)
    if (first_col, first_row, last_col) != (1, HEADERS_ROW, LAST_COLUMN):
        raise ValueError("Unexpected BD Geral table range; stop before exporting")
    if not 1534 <= last_body_row <= ws.max_row:
        raise ValueError("Unexpected BD Geral body size; stop before exporting")
    headers = [text_value(ws.cell(HEADERS_ROW, index).value) for index in range(1, LAST_COLUMN + 1)]
    expected = ("ITEM", "DATA DE FATURAMENTO PREVISTA")
    if headers[0] != expected[0] or headers[15] != expected[1]:
        raise ValueError("Unexpected BD Geral column layout; stop before exporting")
    rows_path = output_dir / "bd_geral_rows.jsonl"
    counts = {"detail_candidate": 0, "other_content_review": 0, "formula_only": 0, "blank": 0}
    error_count = 0
    row_hashes = []
    with rows_path.open("w", encoding="utf-8", newline="\n") as stream:
        for excel_row in range(FIRST_BODY_ROW, last_body_row + 1):
            cells = {}
            constants = 0
            formula_count = 0
            errors = []
            for index in range(1, LAST_COLUMN + 1):
                if index in EXCLUDED_COLUMNS:
                    continue
                letter = get_column_letter(index)
                cell = ws.cell(excel_row, index)
                cached_cell = values_ws.cell(excel_row, index)
                formula = serializable(cell.value) if cell.data_type == "f" else None
                raw_value = None if formula else serializable(cell.value)
                cached_value = serializable(cached_cell.value)
                if formula:
                    formula_count += 1
                elif raw_value not in (None, ""):
                    constants += 1
                if cached_cell.data_type == "e" or (
                    isinstance(cached_value, str) and cached_value.startswith("#")
                ):
                    error_count += 1
                    errors.append(f"{letter}:cached_excel_error")
                cells[letter] = {
                    "header": headers[index - 1],
                    "raw": raw_value,
                    "formula": formula,
                    "cached": cached_value,
                    "numberFormat": cell.number_format,
                    "dataType": cell.data_type,
                }

            item = text_value(ws.cell(excel_row, 1).value)
            service_month = month_value(values_ws.cell(excel_row, 2).value)
            if excel_row in conflict_rows:
                errors.append("conflicted_copy_business_values_differ")
            if constants == 0 and formula_count == 0:
                classification = "blank"
            elif constants == 0:
                classification = "formula_only"
            elif item and service_month:
                classification = "detail_candidate"
            else:
                classification = "other_content_review"
            counts[classification] += 1
            if classification == "other_content_review":
                if not item:
                    errors.append("missing_legacy_item")
                if not service_month:
                    errors.append("missing_service_month")
            row_sha = hashlib.sha256(canonical_json(cells).encode("utf-8")).hexdigest()
            row_hashes.append(row_sha)
            row = {
                "excelRow": excel_row,
                "rowSha256": row_sha,
                "classification": classification,
                "hiddenByFilter": bool(ws.row_dimensions[excel_row].hidden),
                "legacyItem": item,
                "serviceMonth": service_month,
                "engineer": text_value(values_ws.cell(excel_row, 3).value),
                "contractNumber": text_value(values_ws.cell(excel_row, 4).value),
                "clientName": text_value(values_ws.cell(excel_row, 5).value),
                "municipality": text_value(values_ws.cell(excel_row, 6).value),
                "workName": text_value(values_ws.cell(excel_row, 7).value),
                "billingMonth": month_value(values_ws.cell(excel_row, 13).value),
                "expectedReceiptMonth": month_value(values_ws.cell(excel_row, 20).value),
                "receivedMonthText": text_value(values_ws.cell(excel_row, 23).value),
                "cells": cells,
                "issues": errors,
            }
            stream.write(canonical_json(row) + "\n")

    manifest = {
        "formatVersion": 1,
        "sheet": "BD Geral",
        "tableRange": table.ref,
        "sourceWorkbookSha256": source_sha,
        "sourceLabel": "Planilha recebida, provisória; divergências históricas ainda não reconciliadas",
        "capturedAt": dt.datetime.now(dt.timezone.utc).isoformat(),
        "physicalRowCount": last_body_row - FIRST_BODY_ROW + 1,
        "detailCandidateCount": counts["detail_candidate"],
        "includedColumnCount": LAST_COLUMN - len(EXCLUDED_COLUMNS),
        "excludedColumns": ["H", "I"],
        "classifications": counts,
        "cachedErrorCellCount": error_count,
        "knownConflictExcelRows": sorted(conflict_rows),
        "rowsSha256": hashlib.sha256(rows_path.read_bytes()).hexdigest(),
        "rowHashesSha256": hashlib.sha256("\n".join(row_hashes).encode("ascii")).hexdigest(),
        "headers": {get_column_letter(index): headers[index - 1] for index in range(1, LAST_COLUMN + 1) if index not in EXCLUDED_COLUMNS},
        "status": "STAGED",
        "limitations": [
            "Conflicted workbook copy still requires reconciliation",
            "Columns H and I were excluded from the JR import by user decision",
            "Legacy ITEM is not a stable row identity",
            "No operational accounting facts are created by this package",
        ],
    }
    (output_dir / "manifest.json").write_text(canonical_json(manifest) + "\n", encoding="utf-8")
    print(canonical_json({
        "sourceWorkbookSha256": source_sha,
        "rowsSha256": manifest["rowsSha256"],
        "physicalRowCount": manifest["physicalRowCount"],
        "classifications": counts,
        "cachedErrorCellCount": error_count,
    }))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=pathlib.Path)
    parser.add_argument("output_dir", type=pathlib.Path)
    args = parser.parse_args()
    try:
        extract(args.source, args.output_dir)
    except Exception as error:
        print(f"BD Geral export failed: {error}", file=sys.stderr)
        raise
